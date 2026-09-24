"""
针灸文献库 L0 语料治理 + L1 元数据抽取（原型）

用法（工作目录 action-backend/）::

    python -m tools.lit_extract --limit 50 --topic 便秘

设计要点见对话记录，三条硬约束：

1. **NFKC 规范化必须在送模型之前**：99% 的中文文献是全角（``４０例``、``（Ｐ＜０．０５）``），
   不归一化，抽出来的样本量和 p 值也是全角字符串，后面所有数值比较静默失效。
2. **`目录.xlsx` 的标题当权威锚放进 prompt**：实测 #1 的第 1 页夹带了整篇 HIV/AIDS
   研究的正文与数据表，没有锚的话会被抽成一篇 HIV 研究。
3. **年份必须给依据**：实测首 2 页平均出现 4 个不同年份（最多 18 个）——参考文献年、
   基金项目年、收稿年、作者出生年。不要求给依据，模型会挑一个且从不说自己在猜。
"""

from __future__ import annotations

import argparse
import asyncio
import json
import os
import re
import sys
import time
import unicodedata

import fitz
import httpx
import psycopg2
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from module_action.service.publication_parse import parse_ai_json  # noqa: E402
from utils.crypto_util import CryptoUtil  # noqa: E402

CORPUS = r'E:\code\广中医Action\567编号版'
OUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'lit_out')
#: 送进模型的正文上限。中位 17.9k 字符，P90 47.9k；取 30k 覆盖到结果与结论，
#: 尾部的参考文献列表对抽取没有价值，截掉反而省钱
MAX_CHARS = 30000
CONCURRENCY = 6

SCHEMA_PROMPT = """你是针刺临床研究文献的元数据抽取器。下面给你一篇论文的正文（已做全角→半角规范化）。

这篇论文的权威标题是：
《{title}》

严格按这个标题判断哪些内容属于本文。**正文里可能夹带其他论文的片段**（期刊 PDF 跨页所致），
遇到与该标题无关的内容一律忽略，并在 contamination 里说明。

只输出一个 JSON 对象，不要解释、不要 markdown 围栏。字段如下：

{{
  "year": 出版年(整数)或null,
  "year_evidence": "你判断出版年所依据的正文原文片段(原样照抄，20字以内)",
  "year_confidence": "high" 或 "low",
  "authors": ["作者姓名", ...],
  "journal": "期刊名"或null,
  "country": "国家"或null,
  "abstract": "摘要全文(有就照抄，没有就从正文提炼2-4句)",
  "population": "研究对象的具体描述(含年龄/病程等限定)"或null,
  "condition": ["病种"],
  "intervention": ["干预措施"],
  "comparator": ["对照措施"],
  "outcomes": ["结局指标"],
  "tcm_syndrome": ["中医证候"],
  "study_design": "rct|nrct|animal|sr_meta|case_series|other",
  "sample_size": 总样本量(整数)或null,
  "effect_direction": "positive|negative|no_difference|unclear",
  "finding_summary": "主要结果的一句话概括，含关键数据",
  "acupuncture": {{
    "points": ["穴位，优先用国际代码如ST36，无代码则用中文名"],
    "meridian": ["经络"],
    "method": "针法(毫针/电针/温针/火针/耳针/头针/针刀等)"或null,
    "frequency": "电针频率"或null,
    "deqi": true/false/null,
    "retention_min": 留针分钟数(整数)或null,
    "sessions": 总治疗次数(整数)或null,
    "course_weeks": 疗程周数(整数)或null
  }},
  "contamination": true/false,
  "contamination_note": "夹带内容的说明"或null
}}

三条硬约束：
1. **抽不到就填 null 或空数组，绝不推断、绝不编造。** 尤其 sample_size 和针灸参数，
   编一个"n=60""留针30分钟"看着完全正常，是最隐蔽的错误。
2. **year 必须有 year_evidence 支撑。** 正文里通常有多个年份（参考文献年、基金年、
   收稿年、作者出生年）。找不到能确定出版年的依据时，year_confidence 填 "low"。
3. effect_direction 按主要结局判：优于对照=positive，劣于=negative，
   无统计学差异=no_difference，看不出=unclear。

正文：
---
{body}
---"""


def load_catalog() -> dict[int, str]:
    """读 目录.xlsx，返回 编号 → 权威标题"""
    wb = load_workbook(os.path.join(CORPUS, '目录.xlsx'), read_only=True, data_only=True)
    out: dict[int, str] = {}
    for row in wb.worksheets[0].iter_rows(min_row=2, values_only=True):
        if row[0] and row[1]:
            try:
                out[int(row[0])] = str(row[1]).strip()
            except (TypeError, ValueError):
                continue
    wb.close()
    return out


def read_pdf(path: str) -> tuple[str, int]:
    """抽正文并做 NFKC 规范化，返回 (正文, 页数)"""
    d = fitz.open(path)
    text = '\n'.join(d[i].get_text() for i in range(d.page_count))
    pages = d.page_count
    d.close()
    return unicodedata.normalize('NFKC', text), pages


def pick_model() -> tuple[str, str, str, bool]:
    """
    选模型。优先用环境变量（临时试跑用），否则从 `ai_models` 取（生产路径）。

    环境变量：``LLM_MODEL`` / ``LLM_KEY`` / ``LLM_BASE_URL``

    :return: (model_code, api_key, endpoint, 是否下发 thinking 参数)
    """
    if os.environ.get('LLM_KEY') and os.environ.get('LLM_MODEL'):
        code = os.environ['LLM_MODEL']
        key = os.environ['LLM_KEY']
        base = os.environ.get('LLM_BASE_URL', 'https://api.openai.com/v1')
        # `thinking` 是火山方舟/DeepSeek 系的开关，别家不认会 400
        ark = 'volces.com' in base
    else:
        conn = psycopg2.connect(host='127.0.0.1', port=5432, user='postgres', password='123456', dbname='action')
        cur = conn.cursor()
        cur.execute("select model_code, api_key, base_url from ai_models where status='0' order by model_sort limit 1")
        row = cur.fetchone()
        cur.close()
        conn.close()
        if not row:
            raise SystemExit('ai_models 里没有可用模型，且未设置 LLM_KEY/LLM_MODEL 环境变量')
        code, key, base = row
        key = CryptoUtil.decrypt(key)
        ark = 'volces.com' in (base or '')
    url = (base or 'https://api.openai.com/v1').strip().rstrip('/')
    if not url.endswith('/chat/completions'):
        url = url + ('/v3/chat/completions' if url.endswith('/api') else '/chat/completions')
    return code, key, url, ark


async def extract_one(client, model, key, url, no, title, body, sem, ark=False):
    """抽一篇。失败返回 error 字段，不抛。"""
    prompt = SCHEMA_PROMPT.format(title=title, body=body[:MAX_CHARS])
    payload = {'model': model, 'messages': [{'role': 'user', 'content': prompt}],
               'temperature': 0, 'stream': False}
    if ark:
        payload['thinking'] = {'type': 'disabled'}
    async with sem:
        for attempt in (1, 2):
            try:
                r = await client.post(
                    url,
                    headers={'Authorization': f'Bearer {key}', 'Content-Type': 'application/json'},
                    json=payload,
                )
                r.raise_for_status()
                raw = (r.json()['choices'][0]['message']['content'] or '').strip()
                text = raw.strip()
                fence = re.search(r'```(?:json)?\s*(.+?)\s*```', text, re.S)
                if fence:
                    text = fence.group(1)
                s, e = text.find('{'), text.rfind('}')
                data = json.loads(text[s : e + 1])
                data['id'] = no
                data['title'] = title
                return data
            except Exception as exc:  # noqa: BLE001
                if attempt == 2:
                    return {'id': no, 'title': title, 'error': f'{type(exc).__name__}: {exc}'[:200]}
                await asyncio.sleep(2)


async def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--limit', type=int, default=50)
    ap.add_argument('--topic', default='', help='文件名正则，用于挑主题相关的样本')
    ap.add_argument('--others', type=int, default=0, help='额外掺入多少篇无关文献做干扰项')
    ap.add_argument('--ids', default='', help='按编号挑，逗号分隔（优先于 --topic）')
    ap.add_argument('--append', action='store_true', help='追加写入而不是覆盖 papers.jsonl')
    args = ap.parse_args()

    os.makedirs(OUT_DIR, exist_ok=True)
    catalog = load_catalog()
    files = sorted(f for f in os.listdir(CORPUS) if f.lower().endswith('.pdf'))

    if args.ids:
        want = {int(x) for x in args.ids.split(',') if x.strip()}
        picked = [f for f in files if (lambda m: m and int(m.group(1)) in want)(re.match(r'^(\d{1,3})[.\-_ ]', f))]
    elif args.topic:
        hit = [f for f in files if re.search(args.topic, f, re.I)][: args.limit]
        rest = [f for f in files if f not in hit]
        picked = hit + rest[:: max(1, len(rest) // max(args.others, 1))][: args.others]
    else:
        picked = files[: args.limit]
    print(f'选中 {len(picked)} 篇（主题命中 {len([f for f in picked if not args.topic or re.search(args.topic, f, re.I)])}）')

    model, key, url, ark = pick_model()
    print(f'模型 {model} @ {url}\n')

    tasks_in = []
    for name in picked:
        m = re.match(r'^(\d{1,3})[.\-_ ]', name)
        no = int(m.group(1)) if m else None
        title = catalog.get(no, name)
        try:
            body, pages = read_pdf(os.path.join(CORPUS, name))
        except Exception as exc:  # noqa: BLE001
            print(f'  #{no} 解析失败 {exc}')
            continue
        if len(body) < 500:
            print(f'  #{no} 文本过短({len(body)}字符)，疑似扫描件，跳过')
            continue
        tasks_in.append((no, title, body, name, pages))

    sem = asyncio.Semaphore(CONCURRENCY)
    t0 = time.time()
    async with httpx.AsyncClient(timeout=180) as client:
        results = await asyncio.gather(
            *[extract_one(client, model, key, url, no, ti, bo, sem, ark) for no, ti, bo, _, _ in tasks_in]
        )

    meta = {no: (name, pages) for no, _, _, name, pages in tasks_in}
    for r in results:
        name, pages = meta.get(r['id'], ('', 0))
        r['file'] = name
        r['pages'] = pages

    out = os.path.join(OUT_DIR, 'papers.jsonl')
    with open(out, 'a' if args.append else 'w', encoding='utf-8') as fh:
        for r in results:
            fh.write(json.dumps(r, ensure_ascii=False) + '\n')

    ok = [r for r in results if 'error' not in r]
    print(f'\n完成 {len(ok)}/{len(results)}，耗时 {time.time() - t0:.0f}s → {out}')
    for r in results:
        if 'error' in r:
            print(f'  失败 #{r["id"]}: {r["error"]}')


if __name__ == '__main__':
    asyncio.run(main())
