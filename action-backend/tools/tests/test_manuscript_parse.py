"""稿件文件解析的单测（不连库、不连网、不调模型）。

盯的是几件「错了会安静地毁掉整条链路」的事：

- **表格不能漏**：基线特征、结局指标在中文论文里几乎总在表格里。漏掉的话第三步会把
  一整批条目判成「未报告」，而用户明明写了 —— 他会以为是判定错了。
- **换行就是行号**：下游的「报告于第 N 行」、回填草稿、原文面板全以行为锚点
  （定义见 `engine/audit.py::line_map`）。段落没变成换行，行号就没有意义。
- **扫描件要如实报错**：抽不出文字就说抽不出。让用户等几分钟再收获一份全是「未报告」
  的判定，比当场说清楚糟得多。

跑法：python -m pytest tools/tests/test_manuscript_parse.py -q
"""

from __future__ import annotations

import io

import pytest
from module_action.service.manuscript_parse import ManuscriptParseError, parse_manuscript


def _docx(build) -> bytes:
    """按 build(doc) 造一个 docx，返回字节。"""
    from docx import Document

    doc = Document()
    build(doc)
    buf = io.BytesIO()
    doc.save(buf)

    return buf.getvalue()


# --------------------------------------------------------------------------- 纯文本


def test_txt_collapses_blank_runs() -> None:
    """连续空行压成一个 —— docx 里大量空段落很常见，留着会把原文面板撑得没法读。"""
    assert parse_manuscript('a.txt', '第一行\n\n\n\n第二行'.encode()) == '第一行\n\n第二行'


def test_txt_reads_gbk() -> None:
    """中文 Windows 上「另存为 ANSI」的 txt 就是 gbk，而那是国内研究者最容易产出的文件。"""
    assert parse_manuscript('a.txt', '中文内容'.encode('gbk')) == '中文内容'


def test_txt_strips_trailing_whitespace_per_line() -> None:
    # 行尾空白会让 line_map 的 rstrip 与这里不一致，先在这一层统一掉
    assert parse_manuscript('a.md', '甲   \n乙\t\n'.encode()) == '甲\n乙'


# --------------------------------------------------------------------------- docx


def test_docx_keeps_paragraphs_as_lines() -> None:
    def build(doc):
        doc.add_paragraph('引言')
        doc.add_paragraph('本研究纳入 96 例受试者。')

    out = parse_manuscript('a.docx', _docx(build))

    assert out.splitlines() == ['引言', '本研究纳入 96 例受试者。']


def test_docx_extracts_tables() -> None:
    """表格漏了，第三步会把一整批条目误判成「未报告」。"""
    def build(doc):
        doc.add_paragraph('基线特征')
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text = '组别'
        t.cell(0, 1).text = '例数'
        t.cell(1, 0).text = '电针组'
        t.cell(1, 1).text = '48'

    out = parse_manuscript('a.docx', _docx(build))

    assert '组别' in out
    assert '电针组' in out
    assert '48' in out


def test_docx_skips_fully_empty_table_rows() -> None:
    def build(doc):
        doc.add_paragraph('表')
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text = '甲'

    out = parse_manuscript('a.docx', _docx(build))

    # 全空的那一行不该变成一行空白占掉版面
    assert out.splitlines() == ['表', '甲']


def test_doc_is_rejected_with_actionable_advice() -> None:
    """`.doc` 是这里最常见的一次失败，「不支持」三个字帮不上忙。"""
    with pytest.raises(ManuscriptParseError) as got:
        parse_manuscript('稿件.doc', b'\xd0\xcf\x11\xe0')

    assert '另存为 .docx' in str(got.value)


def test_unknown_extension_is_rejected() -> None:
    """白名单而不是黑名单：这一层要把字节喂给解析库，未知格式等于把攻击面暴露出去。"""
    with pytest.raises(ManuscriptParseError):
        parse_manuscript('a.exe', b'MZ\x90\x00')


def test_empty_file_is_rejected() -> None:
    with pytest.raises(ManuscriptParseError):
        parse_manuscript('a.txt', b'')


# --------------------------------------------------------------------------- pdf 重排


def test_pdf_reflow_glues_hard_wrapped_sentences() -> None:
    """PDF 按视觉行断句，不粘回去的话「第 12 行」会指着半句话，回填进草稿的也是半句。"""
    from module_action.service.manuscript_parse import _reflow_pdf

    out = _reflow_pdf('本研究纳入 96 例受试\n者，随机分为两组。\n下一段开始。')

    assert out.splitlines()[0] == '本研究纳入 96 例受试者，随机分为两组。'


def test_pdf_reflow_keeps_new_blocks_apart() -> None:
    """标题、编号、表图说明是新段落的起头，粘上去比不粘更难读。"""
    from module_action.service.manuscript_parse import _reflow_pdf

    out = _reflow_pdf('这一句没有句末标点\n1. 编号项\n表 1 基线特征')

    assert out.splitlines() == ['这一句没有句末标点', '1. 编号项', '表 1 基线特征']


def test_pdf_reflow_spaces_between_latin_but_not_cjk() -> None:
    """中文之间不补空格，中英之间补 —— 否则英文单词会被粘成一坨。"""
    from module_action.service.manuscript_parse import _reflow_pdf

    assert _reflow_pdf('中文一\n中文二。') == '中文一中文二。'
    assert _reflow_pdf('a randomised\ncontrolled trial.') == 'a randomised controlled trial.'
