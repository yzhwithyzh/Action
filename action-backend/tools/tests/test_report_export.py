"""报告草稿导出的单测（不连库、不调模型）。

盯的是几件「错了会直接骗人」的事：
- **docx 正文里绝不能出现条目要求原文** —— 它长得就像一句合格的报告，混进稿件后
  第三步会把这句要求本身判成「已报告」。但 **附录对照表里必须有**：那张表的用途正是
  让编辑逐条对照「这条要求你写在哪」。两者的区别是「以稿件身份」还是「以对照表身份」出现
- 三种格式都得是能打开的真文件（docx/xlsx 是 zip，json 能解析）
- media type 必须显式给，否则前端分不清「一份文件」和「同一条通道回来的错误信封」
"""

from __future__ import annotations

import io
import json
import zipfile

from module_action.service.report_export_service import (
    BUILDERS,
    FORMATS,
    ExportDraft,
    ExportItem,
    build_docx,
    build_json,
    build_xlsx,
)

REQUIREMENT = '文题能识别是随机临床试验'
BODY = '本研究为一项随机对照试验。'

DRAFT = ExportDraft(
    title='电针治疗膝骨关节炎 RCT',
    guideline_code='STRICTA',
    items=(
        ExportItem(1, '1a', '题目和摘要', REQUIREMENT, '0', BODY, 'reported'),
        ExportItem(2, '1b', '题目和摘要', '结构式摘要', '1', '', ''),
        ExportItem(3, '6a', '方法', '针刺细节：选穴依据、进针深度与得气反应', '0', '', 'missing'),
    ),
)


def _docx_text(payload: bytes) -> str:
    """把 docx 拆成纯文本。只读 document.xml，附录表格也在里面"""
    with zipfile.ZipFile(io.BytesIO(payload)) as z:
        return z.read('word/document.xml').decode('utf-8')


def test_all_formats_produce_openable_files() -> None:
    for fmt in FORMATS:
        build, ext, media = BUILDERS[fmt]
        payload = build(DRAFT, 'zh')
        assert payload, fmt
        assert ext == fmt
        assert media, fmt
        if fmt in ('docx', 'xlsx'):
            # OOXML 就是个 zip，能打开说明不是半截文件
            assert zipfile.is_zipfile(io.BytesIO(payload)), fmt


def test_requirement_text_is_in_the_appendix_not_the_prose() -> None:
    """这条是整个导出里最要命的一处，理由见模块 docstring。"""
    xml = _docx_text(build_docx(DRAFT, 'zh'))
    # 用户写的正文在
    assert BODY in xml
    # 要求原文也在 —— 但只该出现在附录表格里，所以它必须在附录标题之后
    assert REQUIREMENT in xml
    assert xml.index('附录') < xml.index(REQUIREMENT), '条目要求出现在了正文部分'


def test_required_gaps_are_marked_optional_gaps_are_not() -> None:
    """缺口要在稿件里看得见，但不能让选填项的占位把初稿淹没。"""
    xml = _docx_text(build_docx(DRAFT, 'zh'))
    assert '【待补充：条目 6a】' in xml
    assert '【待补充：条目 1b】' not in xml


def test_placeholder_carries_no_requirement_text() -> None:
    """占位标记只留条目号 —— 带上要求原文就等于把要求抄进了稿件。"""
    only_blank = ExportDraft(guideline_code='X', items=(ExportItem(1, '1a', '方法', REQUIREMENT, '0', '', ''),))
    xml = _docx_text(build_docx(only_blank, 'zh'))
    body, _, appendix = xml.partition('附录')
    assert '【待补充：条目 1a】' in body
    assert REQUIREMENT not in body
    assert REQUIREMENT in appendix


def test_xlsx_carries_the_verdict_column() -> None:
    """对照表要带上第三步的判定 —— 那正是编辑想看的「你自己查过没有」。"""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(build_xlsx(DRAFT, 'zh')))
    ws = wb.active
    head = [c.value for c in ws[1]]
    assert '校验判定' in head
    col = head.index('校验判定') + 1
    assert ws.cell(row=2, column=col).value == '已报告'
    assert ws.cell(row=4, column=col).value == '未报告'


def test_json_round_trips() -> None:
    payload = json.loads(build_json(DRAFT).decode('utf-8'))
    assert payload['guidelineCode'] == 'STRICTA'
    assert payload['itemTotal'] == 3
    assert payload['filledCount'] == 1
    assert payload['items'][0]['content'] == BODY


def test_english_export_uses_english_labels() -> None:
    xml = _docx_text(build_docx(DRAFT, 'en'))
    assert 'Appendix' in xml
    assert '[To be completed: item 6a]' in xml
    assert '附录' not in xml


def test_untitled_draft_still_exports() -> None:
    """没起名字的草稿照样导得出来 —— 标题回落到「报告初稿」，不该抛异常。"""
    bare = ExportDraft(guideline_code='CARE', items=(ExportItem(9, '1', '引言', '背景', '0', '正文', ''),))
    for fmt in FORMATS:
        assert BUILDERS[fmt][0](bare, 'zh')


# --------------------------------------------------------------------------- 导入来的稿件
#
# 「原稿与条目框两份并存」的兑现点。按条目重拼会把用户的论文打散重排 ——
# 一句话常同时答好几条，而讨论、致谢、参考文献不对应任何条目、会凭空消失。
# 对一份要投出去的论文，这不能接受。

IMPORTED = ExportDraft(
    title='导入的 RCT 稿件',
    guideline_code='STRICTA',
    source_text='# 电针治疗慢性腰痛\n\n## 讨论\n本研究存在若干局限。\n\n## 致谢\n感谢参与本研究的受试者。',
    items=(
        ExportItem(1, '1a', '题目和摘要', REQUIREMENT, '0', '# 电针治疗慢性腰痛', 'reported'),
        ExportItem(3, '6a', '方法', '针刺细节', '0', '', 'missing'),
    ),
)


def test_docx_body_is_the_original_manuscript_when_imported() -> None:
    """导入来的草稿，正文必须是原稿本身。

    **不对应任何条目的段落也要在**（讨论、致谢）—— 它们正是按条目重拼时会丢掉的那些，
    也是用户最不能接受丢掉的那些。
    """
    text = _docx_text(build_docx(IMPORTED, 'zh'))

    assert '本研究存在若干局限。' in text
    assert '感谢参与本研究的受试者。' in text


def test_docx_keeps_the_appendix_table_for_imported_drafts() -> None:
    """正文换成原稿，但条目对照表照旧 —— 投稿时随稿提交的正是那张表。"""
    text = _docx_text(build_docx(IMPORTED, 'zh'))

    assert REQUIREMENT in text


def test_docx_omits_placeholders_when_exporting_the_original() -> None:
    """原稿是用户写好的成品，不该被塞进「【待补充：6a】」这类占位符。

    占位符是「逐条填」那条路的产物：那边缺口要在稿件里看得见。这边缺口的正确去处是
    附录对照表的「未填写」列，往人家的论文正文里插占位符是在改稿。
    """
    text = _docx_text(build_docx(IMPORTED, 'zh'))

    assert '待补充' not in text
    # 逐条填的草稿仍然要有占位符
    assert '待补充' in _docx_text(build_docx(DRAFT, 'zh'))
