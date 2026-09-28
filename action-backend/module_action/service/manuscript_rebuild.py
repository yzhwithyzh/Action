"""
「生成初稿」：把条目上的修改写回原稿（纯函数，可离线单测，见 `tools/tests/test_manuscript_rebuild.py`）。

## 要解决的问题

导入稿件时，稿件里的段落按判定行号回填进各条目（`ChecklistReviewService._backfill_draft`），
之后用户在条目上手改、用「AI 完善」扩写。可是**导出的正文一直是导入时的原稿** —— 条目上的
改动只出现在附录对照表里，改了等于没改。反过来按条目重拼全文也不行：讨论、致谢、参考文献
不对应任何条目，图片、表格、样式更是一样都带不过来。

所以这里做的是**在原稿上就地替换**：原稿（原文件字节）原封不动地留着，每次生成都从它出发，
把「与导入时相比变了」的条目写回它原来所在的段落，其余一概不动。图片、表格、页眉页脚、
样式、没对应条目的章节全部保留。**每次都从原稿重算**，所以生成多少次结果都一样，
也不存在「上一次生成的稿子被当成了原稿」这种越改越乱的情况。

## 规则（按条目在清单里的顺序逐条处理，只处理「变了」的条目）

「变了」= 当前正文与导入时回填的原文去掉空白后不相等。没变的条目一个字都不碰。

1. **原位替换**：这一条导入时来自哪几段，就替换哪几段 —— 第一段换成新正文（新正文有几段
   就拆成几段，沿用原段落的样式），其余几段删掉。段落里嵌着的图片保留，只换文字。
2. **共用的段落不替换，改为追加**：一句话常同时回答好几条。那段话若也属于某个**没改**的条目，
   替换掉就等于把那一条的内容删了 —— 所以保留原段落，新正文插在它后面。统计里单独计数，
   前台提示用户核对有没有重复。
3. **表格不改**：表格里的一行被改成一段文字会毁掉整张表。新正文插在表格后面，同样计入「追加」。
4. **导入时没找到位置的条目**（用户后来才写的）：插在清单顺序上最近的**前一个**有位置的条目
   之后；前面都没有就插在最近的**后一个**之前；再没有就放到全文末尾。清单顺序与论文章节
   顺序基本一致（题目 → 摘要 → 引言 → 方法 → 结果 → 讨论），这样插进去的段落落在对的章节里。
5. **清空了的条目**：删掉它独占的原段落（共用的、带图片的段落照旧保留）。

## 位置按「文字」对，不只按「行号」对

回填时记下的是行号。行号 → 那一行的原文 → 在原稿的块里找同样的文字（取序号最接近的那处）。
只按行号对的话，解析规则一变（例如 docx 的表格从「统一追加在末尾」改成「按原位」）
老草稿的行号就全部错位，而错位的表现是「改动写进了不相干的段落」，极难察觉。

## 输出

- docx 原稿：直接改原文档，样式、图片、表格、页面设置全在
- PDF / 纯文本 / 没存原文件的老草稿：按块新建一份 docx（PDF 里的图片按原位置插回，
  表格只剩文字 —— PDF 本来就没有表格结构）
- 纯文本：与 docx 同序，图片处写一行「［图 n］」，给预览与一致性检查用
"""

from __future__ import annotations

import copy
import io
import re
from dataclasses import dataclass, field
from typing import Any

from module_action.service.manuscript_parse import (
    Block,
    docx_blocks,
    ext_of,
    open_docx,
    pdf_blocks,
    text_blocks,
)

W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
_W = f'{{{W_NS}}}'
#: 插进来的图片最宽多少（厘米）—— A4 版心约 15–16 cm
_MAX_IMAGE_CM = 15.0
#: 图片占位行，预览与一致性检查里代表一张图
IMAGE_MARK = {'zh': '［图 {n}］', 'en': '[Figure {n}]'}


@dataclass(frozen=True)
class ItemEdit:
    """
    一条 checklist 条目的「导入时 vs 现在」

    :ivar item_id: 条目id
    :ivar lines: 导入时这一条回填自原稿的行号（空 = 导入时没找到对应段落）
    :ivar baseline: 导入时回填进去的原文
    :ivar current: 现在的正文
    """

    item_id: int
    lines: tuple[int, ...] = ()
    baseline: str = ''
    current: str = ''


@dataclass
class RebuildStats:
    """这次生成改了什么，前台据此给一句说明"""

    #: 原位替换的条目数
    replaced: int = 0
    #: 新插入的条目数（导入时没找到位置）
    inserted: int = 0
    #: 因与未改动条目共用段落、或位于表格中，没替换而是追加在原段落之后的条目数
    appended: int = 0
    #: 被清空、原段落已删除的条目数
    removed: int = 0
    #: 生成后的稿件里有几张图、几个表格
    images: int = 0
    tables: int = 0


@dataclass
class RebuildResult:
    text: str
    docx: bytes
    stats: RebuildStats = field(default_factory=RebuildStats)


def _squash(s: str) -> str:
    return re.sub(r'\s+', '', s or '')


def _paras(text: str) -> list[str]:
    """新正文拆成段：一行一段，空行不要"""
    return [ln.strip() for ln in (text or '').splitlines() if ln.strip()]


# ---------------------------------------------------------------- docx 元素操作


def _xp(expr: str) -> Any:
    # python-docx 的元素类改写了 xpath() 的签名，编译好的 lxml XPath 对任何元素都能用
    from lxml import etree  # noqa: PLC0415

    return etree.XPath(expr, namespaces={'w': W_NS})


_MEDIA = _xp('.//w:drawing|.//w:pict|.//w:object')
_PICTURES = _xp('.//w:drawing|.//w:pict')


def _has_media(el: Any) -> bool:
    return bool(_MEDIA(el))


def _first_rpr(p_el: Any) -> Any:
    """段落里第一个文字 run 的格式（没有就 None），新文字沿用它 —— 字体字号与原段落一致"""
    for r in p_el.iter(f'{_W}r'):
        if _has_media(r):
            continue
        rpr = r.find(f'{_W}rPr')
        return copy.deepcopy(rpr) if rpr is not None else None
    return None


def _make_run(text: str, rpr: Any) -> Any:
    # 用 python-docx 的 OxmlElement 而不是裸 lxml：后者造出来的元素不是 CT_P/CT_R，
    # 生成完再用 python-docx 读回这份文档时，新插的段落会读出 None
    from docx.oxml import OxmlElement  # noqa: PLC0415

    r = OxmlElement('w:r')
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    t = OxmlElement('w:t')
    r.append(t)
    t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    return r


#: 段落里这几类子元素是「位置标记」不是文字，换文字时留着
_KEEP_TAGS = {f'{_W}{t}' for t in ('pPr', 'bookmarkStart', 'bookmarkEnd', 'commentRangeStart', 'commentRangeEnd')}


def _strip_text(p_el: Any) -> None:
    """删掉段落里所有不含图片的文字内容（run、超链接、修订标记……），图片 run 留着"""
    for child in list(p_el):
        if child.tag in _KEEP_TAGS or _has_media(child):
            continue
        p_el.remove(child)


def _set_text(p_el: Any, text: str) -> None:
    """把段落文字换成 text：沿用原格式，图片不动"""
    rpr = _first_rpr(p_el)
    _strip_text(p_el)
    ppr = p_el.find(f'{_W}pPr')
    p_el.insert(1 if ppr is not None else 0, _make_run(text, rpr))


def _new_para_like(ref: Any, text: str, heading_ids: set[str]) -> Any:
    """
    照着 ref 的样子新建一段

    ref 是标题时**不照抄样式**（正文段落插在标题后面还是标题样式就乱了），
    是列表项时去掉编号（不然新段落会接着编号），分节符一律去掉（照抄会凭空多出一个分节）
    """
    from docx.oxml import OxmlElement  # noqa: PLC0415

    p = OxmlElement('w:p')
    rpr = None
    if ref is not None and ref.tag == f'{_W}p':
        ppr = ref.find(f'{_W}pPr')
        style = ppr.find(f'{_W}pStyle') if ppr is not None else None
        is_heading = style is not None and style.get(f'{_W}val') in heading_ids
        if ppr is not None and not is_heading:
            ppr = copy.deepcopy(ppr)
            for tag in ('numPr', 'sectPr'):
                for x in ppr.findall(f'{_W}{tag}'):
                    ppr.remove(x)
            p.append(ppr)
        if not is_heading:
            rpr = _first_rpr(ref)
    p.append(_make_run(text, rpr))
    return p


def _heading_ids(doc: Any) -> set[str]:
    """标题类样式的 styleId。中文版 Word 的标题样式 id 是「1」「2」这种，只能按样式名认"""
    out = set()
    for st in doc.styles:
        name = (getattr(st, 'name', '') or '').lower()
        if name.startswith(('heading', 'title', 'subtitle')) or '标题' in name:
            out.add(st.style_id)
    return out


def _new_doc_from_blocks(blocks: list[Block]) -> tuple[Any, list[Block]]:
    """PDF / 纯文本 → 新建一份 docx，块与新文档的段落一一对应"""
    from docx import Document  # noqa: PLC0415
    from docx.shared import Cm, Pt  # noqa: PLC0415

    doc = Document()
    out: list[Block] = []
    for b in blocks:
        if b.kind == 'image':
            p = doc.add_paragraph()
            try:
                width = min(_MAX_IMAGE_CM, max(b.width, 1) / 72 * 2.54)
                p.add_run().add_picture(io.BytesIO(b.image), width=Cm(width))
            except Exception:  # 个别 PDF 内嵌图格式 python-docx 不认，跳过这张图不该让整次生成失败
                continue
            out.append(Block('para', [], el=p._p))
        else:
            p = doc.add_paragraph('\n'.join(b.lines))
            out.append(Block('para', list(b.lines), el=p._p))
    # 默认模板的正文字号偏大，统一压到 11pt
    doc.styles['Normal'].font.size = Pt(11)
    return doc, out


# ---------------------------------------------------------------- 主流程


def _resolve(blocks: list[Block], line_texts: dict[int, str], nums: tuple[int, ...]) -> list[tuple[int, int]]:
    """
    行号 → 原稿里的位置 (块下标, 块内行下标)

    按文字对：行号 n 的原文在块里出现过的几处中，取「序号」最接近 n 的那处。
    """
    positions = [(bi, si, ln) for bi, b in enumerate(blocks) for si, ln in enumerate(b.lines)]
    by_text: dict[str, list[int]] = {}
    for k, (_bi, _si, ln) in enumerate(positions):
        by_text.setdefault(_squash(ln), []).append(k)
    out = []
    for n in sorted(set(nums)):
        text = line_texts.get(n)
        if text is None:
            continue
        cands = by_text.get(_squash(text))
        if not cands:
            continue
        k = min(cands, key=lambda c: abs(c - (n - 1)))
        out.append(positions[k][:2])
    return out


class _Rebuilder:
    """
    一次重建的工作台：持有文档、块、归属关系与插入游标，按条目逐条处理

    拆成类只为把「怎么插」「怎么删」「一条该怎么处理」分开写；规则见模块 docstring。
    """

    def __init__(self, doc: Any, blocks: list[Block], line_texts: dict[int, str], edits: list[ItemEdit]) -> None:
        self.doc = doc
        self.blocks = blocks
        self.edits = edits
        self.heading_ids = _heading_ids(doc)
        self.stats = RebuildStats()
        # 每条在原稿里的位置；每个位置归谁
        self.where = {e.item_id: _resolve(blocks, line_texts, e.lines) for e in edits}
        self.changed = {e.item_id for e in edits if _squash(e.current) != _squash(e.baseline)}
        self.owners: dict[tuple[int, int], set[int]] = {}
        for item_id, pos in self.where.items():
            for p in pos:
                self.owners.setdefault(p, set()).add(item_id)
        #: 已被替换过的块 —— 一块只能被一条替换
        self.claimed: set[int] = set()
        #: 待删的段落。**最后统一删**：先删的话，后面的条目拿它当锚点插入就找不到位置了
        self.to_delete: list[Any] = []
        # 同一个锚点后面连续插好几段时，要接在上一次插入的后面，顺序才对
        self.cursor_after: dict[int, Any] = {}
        self.cursor_before: dict[int, Any] = {}

    # ---- 插入

    def insert_after(self, bi: int, texts: list[str]) -> None:
        ref = self.cursor_after.get(bi, self.blocks[bi].el)
        for t in texts:
            p = _new_para_like(self.blocks[bi].el, t, self.heading_ids)
            ref.addnext(p)
            ref = p
        self.cursor_after[bi] = ref

    def insert_before(self, bi: int, texts: list[str]) -> None:
        anchor = self.blocks[bi].el
        ref = self.cursor_before.get(bi)
        for t in texts:
            p = _new_para_like(anchor, t, self.heading_ids)
            if ref is None:
                anchor.addprevious(p)
            else:
                ref.addnext(p)
            ref = p
        self.cursor_before[bi] = ref

    def append_end(self, texts: list[str]) -> None:
        body = self.doc.element.body
        sect = body.find(f'{_W}sectPr')
        for t in texts:
            p = _new_para_like(None, t, self.heading_ids)
            if sect is not None:
                sect.addprevious(p)
            else:
                body.append(p)

    def drop(self, bi: int) -> None:
        """删一块：带图的段落只清文字（图留着），其余整段删"""
        el = self.blocks[bi].el
        if _has_media(el):
            _strip_text(el)
        else:
            self.to_delete.append(el)

    # ---- 逐条

    def exclusive(self, bi: int, item_id: int) -> bool:
        """这一块能不能由 item_id 替换：是段落、没被替换过、每一行都只归改过的条目且包括它"""
        b = self.blocks[bi]
        if b.kind != 'para' or bi in self.claimed:
            return False
        for si in range(len(b.lines)):
            who = self.owners.get((bi, si), set())
            if item_id not in who or any(o not in self.changed for o in who):
                return False
        return True

    def place_new(self, idx: int, texts: list[str]) -> None:
        """导入时没找到位置的条目：前一个有位置的条目之后 → 后一个之前 → 文末"""
        order = [e.item_id for e in self.edits]
        prev = next((self.where[i] for i in reversed(order[:idx]) if self.where[i]), None)
        if prev:
            self.insert_after(max(bi for bi, _ in prev), texts)
            return
        nxt = next((self.where[i] for i in order[idx + 1:] if self.where[i]), None)
        if nxt:
            self.insert_before(min(bi for bi, _ in nxt), texts)
        else:
            self.append_end(texts)

    def apply(self, idx: int, e: ItemEdit) -> None:
        """处理一条改过的条目"""
        texts = _paras(e.current)
        pos = self.where[e.item_id]
        if not pos:
            if texts:
                self.stats.inserted += 1
                self.place_new(idx, texts)
            return

        my_blocks = sorted({bi for bi, _si in pos})
        mine = [bi for bi in my_blocks if self.exclusive(bi, e.item_id)]
        self.claimed.update(mine)

        if not texts:
            # 清空了：删它独占的段落
            for bi in mine:
                self.drop(bi)
            if mine:
                self.stats.removed += 1
            return

        if not mine:
            # 全是共用段落或表格：原文留着，新正文追加在它最后一块后面
            self.insert_after(my_blocks[-1], texts)
            self.stats.appended += 1
            return

        _set_text(self.blocks[mine[0]].el, texts[0])
        if len(texts) > 1:
            self.insert_after(mine[0], texts[1:])
        for bi in mine[1:]:
            self.drop(bi)
        self.stats.replaced += 1

    def run(self) -> None:
        for idx, e in enumerate(self.edits):
            if e.item_id in self.changed:
                self.apply(idx, e)
        for el in self.to_delete:
            parent = el.getparent()
            if parent is not None:
                parent.remove(el)

    def output(self, locale: str) -> RebuildResult:
        """纯文本按新文档重新读一遍，图片处写占位行"""
        mark = IMAGE_MARK['en' if locale == 'en' else 'zh']
        lines: list[str] = []
        for b in docx_blocks(self.doc):
            if b.kind == 'table':
                self.stats.tables += 1
            lines.extend(b.lines)
            if b.kind != 'para':
                continue
            for _ in _PICTURES(b.el):
                self.stats.images += 1
                lines.append(mark.format(n=self.stats.images))
        buf = io.BytesIO()
        self.doc.save(buf)

        return RebuildResult(text='\n'.join(lines), docx=buf.getvalue(), stats=self.stats)


def rebuild(
    filename: str,
    data: bytes | None,
    source_text: str,
    line_texts: dict[int, str],
    edits: list[ItemEdit],
    locale: str = 'zh',
) -> RebuildResult:
    """
    从原稿出发，把变了的条目写回原处，生成新的稿件

    :param filename: 原稿文件名（取扩展名）；老草稿没存原文件时可为空
    :param data: 原稿字节；None 时退回按 source_text 重建（没有图片与表格可保留）
    :param source_text: 导入时解析出的纯文本（回填行号就是按它编的）
    :param line_texts: {行号: 该行原文}，由调用方按 `engine/audit.py::line_map` 算好传进来
    :param edits: 全部条目，**按清单顺序**
    :param locale: 图片占位行的语言
    :return: 新稿件的纯文本、docx 字节与改动统计
    """
    ext = ext_of(filename or '')
    if data and ext == '.docx':
        doc = open_docx(data)
        blocks = docx_blocks(doc)
    else:
        src_blocks = pdf_blocks(data) if data and ext == '.pdf' else text_blocks(source_text)
        doc, blocks = _new_doc_from_blocks(src_blocks)

    work = _Rebuilder(doc, blocks, line_texts, edits)
    work.run()

    return work.output(locale)
